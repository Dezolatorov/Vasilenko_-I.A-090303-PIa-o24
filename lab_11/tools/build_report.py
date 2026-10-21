"""Сборка отчёта по лабораторной работе №11 в формате .docx."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(r"C:\elkin\lab_11")
SCREENS = BASE / "screens"
OUT = Path(r"C:\elkin\Отчёты\Практика 11 Василенко И.А.docx")

BODY_FONT = "Times New Roman"
MONO_FONT = "Consolas"


def set_style(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(14)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    for style_name, size, bold in (("Heading 1", 14, True), ("Heading 2", 14, True)):
        style = document.styles[style_name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.line_spacing = 1.5


def body(document: Document, text: str, first_indent: float = 1.25, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(first_indent)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = align
    run = paragraph.add_run(text)
    run.font.name = BODY_FONT
    run.font.size = Pt(14)
    return paragraph


def bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(1.0)
    paragraph.paragraph_format.first_line_indent = Cm(-0.5)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = paragraph.add_run("– " + text)
    run.font.name = BODY_FONT
    run.font.size = Pt(14)


def heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = BODY_FONT
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)


def shading(paragraph, color: str = "F2F2F2") -> None:
    properties = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    properties.append(shd)


def listing(document: Document, code: str, caption: str | None = None) -> None:
    if caption:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(caption)
        run.font.name = BODY_FONT
        run.font.size = Pt(12)
        run.font.italic = True

    for line in code.strip("\n").split("\n"):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.left_indent = Cm(0.6)
        shading(paragraph)
        run = paragraph.add_run(line if line else " ")
        run.font.name = MONO_FONT
        run.font.size = Pt(9)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def table_caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.font.name = BODY_FONT
    run.font.size = Pt(12)
    run.font.italic = True


def table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None,
          font_size: int = 11) -> None:
    prepared = document.add_table(rows=1, cols=len(headers))
    prepared.style = "Table Grid"
    prepared.alignment = WD_TABLE_ALIGNMENT.CENTER

    header_cells = prepared.rows[0].cells
    for index, text in enumerate(headers):
        header_cells[index].text = ""
        paragraph = header_cells[index].paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.font.name = BODY_FONT
        run.font.size = Pt(font_size)
        run.font.bold = True

    for row in rows:
        cells = prepared.add_row().cells
        for index, text in enumerate(row):
            cells[index].text = ""
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(str(text))
            run.font.name = BODY_FONT
            run.font.size = Pt(font_size)

    if widths:
        for row in prepared.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)

    document.add_paragraph().paragraph_format.space_after = Pt(0)


def figure(document: Document, path: Path, caption: str, width_cm: float = 15.5) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run().add_picture(str(path), width=Cm(width_cm))

    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(10)
    run = caption_paragraph.add_run(caption)
    run.font.name = BODY_FONT
    run.font.size = Pt(12)


def build() -> Document:
    document = Document()
    set_style(document)

    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)

    header = [
        "Выполнил: Василенко И.А Группа: ПИа-090303-о24",
        "ЛАБОРАТОРНАЯ РАБОТА №11",
        "Дисциплина: Основы программирования на Python Тема: Реляционные и нереляционные "
        "базы данных. Хранение данных и использование CDN",
        "Вариант 2: Библиотека",
    ]
    for line in header:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(line)
        run.font.name = BODY_FONT
        run.font.size = Pt(16)
        run.font.bold = True

    document.add_paragraph()

    heading(document, "Цель работы")
    body(document, "Получить практические навыки проектирования и использования систем хранения данных "
                   "в Python-приложениях: изучить принципы работы реляционных баз данных, освоить SQL, "
                   "первичные и внешние ключи, связи между таблицами и нормализацию; разобрать операции "
                   "SELECT, INSERT, UPDATE, DELETE, JOIN, группировку и агрегатные функции, транзакции; "
                   "подключить Python-приложение к реляционной БД и организовать слой Repository; сравнить "
                   "реляционный и нереляционный подходы; реализовать дополнительное хранилище на базе Redis "
                   "и организовать кэширование с TTL и инвалидацией; изучить назначение CDN и организовать "
                   "раздачу статических ресурсов через CDN.")

    heading(document, "Постановка индивидуального задания (Вариант 2. Библиотека)")
    body(document, "Предметная область — библиотека, основная сущность — Book (книга). Требуется "
                   "модернизировать REST-приложение из лабораторной работы №10: данные о книгах должны "
                   "храниться не в list, а в реляционной базе данных; обращения к БД выполняются через "
                   "слой Repository; NoSQL-часть реализуется на Redis (история поиска и кэш); "
                   "CDN-часть — обложки книг и статические файлы оформления.")
    body(document, "В соответствии с заданием необходимо:")
    for item in (
        "спроектировать схему БД: не менее 3 таблиц, первичный ключ в каждой таблице, не менее 2 внешних "
        "ключей, связь 1:N и связь N:M через промежуточную таблицу;",
        "создать SQL-скрипт schema.sql и заполнить БД тестовыми данными (не менее 10 записей основной "
        "сущности, 5 записей связанных сущностей, записи для связи N:M);",
        "написать не менее 15 SQL-запросов, включая SELECT, WHERE, ORDER BY, JOIN, LEFT JOIN, GROUP BY, "
        "HAVING, COUNT, AVG/SUM, INSERT, UPDATE, DELETE, подзапрос, запрос с несколькими условиями и "
        "запрос со связью нескольких таблиц;",
        "определить не менее двух индексов и объяснить их назначение;",
        "перевести все endpoints лабораторной работы №10 на работу с БД через Repository, используя "
        "только параметризованные запросы;",
        "реализовать бизнес-операцию, требующую транзакции (выдача и возврат книги);",
        "реализовать Redis-кэш по паттерну cache-aside с TTL и инвалидацией, а также историю поиска;",
        "разместить не менее трёх статических ресурсов и получить их через CDN, исследовав заголовки "
        "и поведение кэша;",
        "написать не менее 10 автоматических тестов на pytest.",
    ):
        bullet(document, item)

    heading(document, "Краткие теоретические сведения")
    body(document, "Реляционная база данных представляет данные в виде таблиц, связанных между собой "
                   "ключами. Первичный ключ (PRIMARY KEY) однозначно идентифицирует строку таблицы, внешний "
                   "ключ (FOREIGN KEY) ссылается на первичный ключ другой таблицы и обеспечивает ссылочную "
                   "целостность. Основные типы связей: один к одному (1:1), один ко многим (1:N) и многие ко "
                   "многим (N:M). Связь N:M в реляционной модели реализуется через промежуточную таблицу, "
                   "содержащую внешние ключи обеих сторон.")
    body(document, "Нормализация уменьшает дублирование данных и устраняет аномалии вставки, обновления и "
                   "удаления. Первая нормальная форма требует атомарности значений в ячейках, вторая — "
                   "зависимости неключевых атрибутов от полного первичного ключа, третья — отсутствия "
                   "транзитивных зависимостей между неключевыми атрибутами. В спроектированной схеме сведения "
                   "об авторе вынесены в отдельную таблицу authors, а теги — в tags и book_tags, что "
                   "соответствует 3НФ.")
    body(document, "SQL — декларативный язык работы с реляционными данными: SELECT выбирает строки, INSERT "
                   "добавляет, UPDATE изменяет, DELETE удаляет. Оператор JOIN соединяет таблицы по условию, "
                   "GROUP BY группирует строки, агрегатные функции COUNT, SUM, AVG, MIN, MAX вычисляют "
                   "статистику, HAVING фильтрует уже сгруппированные данные, а подзапрос позволяет "
                   "использовать результат одного запроса внутри другого.")
    body(document, "Транзакция — набор операций, выполняемый как единое целое. Свойства ACID: атомарность "
                   "(все операции выполняются или ни одна), согласованность (БД переходит из одного "
                   "корректного состояния в другое), изолированность (параллельные транзакции не мешают друг "
                   "другу) и долговечность (после COMMIT изменения сохраняются). При ошибке выполняется "
                   "ROLLBACK, и данные возвращаются к состоянию до начала транзакции.")
    body(document, "Индекс — дополнительная структура ускорения поиска, которая хранит упорядоченные "
                   "значения столбца со ссылками на строки. Индекс ускоряет SELECT, но замедляет INSERT, "
                   "UPDATE и DELETE и занимает место на диске, поэтому его создают под конкретные запросы, "
                   "а не на каждый столбец.")
    body(document, "SQL-инъекция — атака, при которой пользовательские данные попадают в текст SQL-запроса "
                   "и меняют его смысл. Приём вида \u00abSELECT * FROM users WHERE name = '\u0027 + name + "
                   "\u0027'\u00bb опасен: значение \"' OR '1'='1\" превращает условие в всегда истинное. "
                   "Параметризованные запросы (placeholders \u0027?\u0027) передают значения отдельно от "
                   "текста запроса, поэтому данные не могут стать частью SQL-кода.")
    body(document, "NoSQL объединяет хранилища, не использующие классическую реляционную модель: документные "
                   "(MongoDB), ключ-значение (Redis), колоночные (Cassandra), графовые (Neo4j). Redis хранит "
                   "данные в памяти в виде ключ-значение, поддерживает TTL и атомарные операции INCR, поэтому "
                   "часто применяется как кэш, хранилище сессий и счётчиков.")
    body(document, "Кэширование по паттерну cache-aside: приложение сначала проверяет кэш, при промахе "
                   "читает данные из основной БД, записывает их в кэш с TTL и возвращает клиенту. TTL "
                   "ограничивает время жизни записи, а cache invalidation удаляет или обновляет запись при "
                   "изменении данных, чтобы клиент не получил устаревшее значение.")
    body(document, "CDN (Content Delivery Network) — распределённая сеть edge-серверов, которые отдают "
                   "статический контент ближе к пользователю и снижают нагрузку на origin-сервер. "
                   "Кэшированием на edge управляет заголовок Cache-Control (например, «public, max-age=86400»), "
                   "а признаками попадания в кэш служат заголовки X-Cache и Age. Через CDN передают "
                   "изображения, CSS, JavaScript, видео и шрифты; изменяемые данные приложения (цены, "
                   "остатки) должны поступать из БД.")

    heading(document, "Проектирование базы данных")
    body(document, "Для предметной области «Библиотека» спроектирована схема из шести таблиц: authors, "
                   "books, readers, loans, tags и book_tags. Схема содержит пять внешних ключей, связь 1:N "
                   "(authors → books, readers → loans, books → loans) и связь N:M (books ↔ tags через "
                   "промежуточную таблицу book_tags).")
    listing(document, """AUTHORS 1 ──── N BOOKS N ──── M TAGS
   id PK            id PK           id PK
   full_name        title           name
   country          author_id FK ─┐
                    genre         │  BOOK_TAGS (промежуточная таблица)
                    year          │    book_id FK ──► BOOKS.id
                    price         │    tag_id  FK ──► TAGS.id
                    total_copies  │    PRIMARY KEY (book_id, tag_id)
                    available     │
                    cover_path    │
                                  │
READERS 1 ──── N LOANS ───────────┘
   id PK           id PK
   full_name       book_id FK  ──► BOOKS.id
   email UNIQUE    reader_id FK ──► READERS.id
   phone                              taken_at
   registered_at   due_at
                   returned_at""", "Схема 11.1 — ER-диаграмма базы данных «Библиотека»")

    body(document, "Назначение таблиц и ограничения целостности приведены в таблице 11.1.")
    table_caption(document, "Таблица 11.1 — Назначение таблиц и ограничения целостности")
    table(document,
          ["Таблица", "Назначение", "Ключи и ограничения"],
          [
              ["authors", "Справочник авторов", "id PK; full_name UNIQUE, NOT NULL"],
              ["books", "Библиотечный фонд", "id PK; author_id FK → authors(id); CHECK (year 1…2100), "
                                            "CHECK (price > 0), CHECK (available_copies ≤ total_copies)"],
              ["readers", "Читатели", "id PK; email UNIQUE, NOT NULL"],
              ["loans", "Выдачи книг", "id PK; book_id FK → books(id) ON DELETE CASCADE; "
                                       "reader_id FK → readers(id); CHECK (due_at ≥ taken_at)"],
              ["tags", "Теги произведений", "id PK; name UNIQUE"],
              ["book_tags", "Связь N:M books ↔ tags", "PRIMARY KEY (book_id, tag_id); два внешних ключа"],
          ],
          widths=[3.0, 5.0, 9.0])

    body(document, "Схема обоснована с точки зрения нормализации: ФИО автора хранится один раз в таблице "
                   "authors (устранено дублирование в books), теги вынесены в отдельный справочник, а связь "
                   "книги и тега описана в промежуточной таблице. Все неключевые атрибуты зависят только от "
                   "первичного ключа своей таблицы, транзитивных зависимостей нет — схема находится в 3НФ.")
    body(document, "База данных создаётся скриптом schema.sql; тестовые данные (12 книг, 6 авторов, "
                   "5 читателей, 8 выдач, 8 тегов, 24 связи N:M) загружаются из seed.sql. Структура "
                   "созданной БД показана на рисунке 11.1.")
    figure(document, SCREENS / "11-01-schema.png",
           "Рисунок 11.1 — Структура базы данных: таблицы, ключи, количество записей и индексы")

    heading(document, "SQL-запросы")
    body(document, "В файле queries.sql подготовлено 15 запросов, охватывающих все обязательные конструкции "
                   "задания. Перечень запросов приведён в таблице 11.2.")
    table_caption(document, "Таблица 11.2 — Перечень SQL-запросов и использованных конструкций")
    table(document,
          ["№", "Конструкция", "Назначение запроса"],
          [
              ["1", "SELECT", "выборка всех книг фонда"],
              ["2", "WHERE", "книги дешевле 600 рублей, кроме антиутопий"],
              ["3", "ORDER BY", "сортировка фонда по цене по убыванию"],
              ["4", "JOIN", "книги вместе с авторами (связь 1:N)"],
              ["5", "LEFT JOIN + COUNT", "читатели и число взятых книг, включая читателей без выдач"],
              ["6", "GROUP BY", "количество книг по жанрам"],
              ["7", "HAVING", "авторы, у которых больше двух книг в фонде"],
              ["8", "COUNT", "количество невозвращённых выдач"],
              ["9", "AVG, SUM", "средняя цена книги и стоимость фонда"],
              ["10", "подзапрос", "книги дороже средней цены фонда"],
              ["11", "несколько условий", "активные выдачи с просрочкой на заданную дату"],
              ["12", "JOIN трёх таблиц", "книги, авторы и теги через book_tags"],
              ["13", "INSERT", "регистрация нового читателя"],
              ["14", "UPDATE", "повышение цены книг жанра «антиутопия» на 10 %"],
              ["15", "DELETE + подзапрос", "удаление читателя, не бравшего книг"],
          ],
          widths=[1.2, 4.3, 11.5])

    listing(document, """-- 4. JOIN: книги вместе с авторами (связь 1:N)
SELECT b.title, a.full_name, a.country
FROM books AS b
JOIN authors AS a ON a.id = b.author_id
ORDER BY a.full_name, b.title;

-- 7. HAVING: авторы, у которых больше двух книг в фонде
SELECT a.full_name, COUNT(b.id) AS books_count
FROM authors AS a
JOIN books AS b ON b.author_id = a.id
GROUP BY a.id, a.full_name
HAVING COUNT(b.id) > 2
ORDER BY books_count DESC;

-- 12. Связь нескольких таблиц через промежуточную таблицу
SELECT b.title, a.full_name, t.name AS tag
FROM books AS b
JOIN authors AS a    ON a.id = b.author_id
JOIN book_tags AS bt ON bt.book_id = b.id
JOIN tags AS t       ON t.id = bt.tag_id
WHERE t.name = 'фантастика'
ORDER BY b.title;""", "Листинг 11.1 — Примеры запросов из queries.sql")

    body(document, "Результаты выполнения запросов приведены на рисунках 11.2 и 11.3.")
    figure(document, SCREENS / "11-02-sql.png", "Рисунок 11.2 — Выполнение SQL-запросов № 1–6")
    figure(document, SCREENS / "11-03-sql.png", "Рисунок 11.3 — Выполнение SQL-запросов № 7–15")

    heading(document, "Индексы")
    body(document, "В схеме создано пять индексов. Каждый индекс предназначен для конкретного запроса "
                   "приложения, а не «на всякий случай» (таблица 11.3).")
    table_caption(document, "Таблица 11.3 — Индексы базы данных и их назначение")
    table(document,
          ["Индекс", "Столбцы", "Для какого запроса"],
          [
              ["idx_books_title", "books(title)", "поиск и сортировка по названию: WHERE title = ? / ORDER BY title"],
              ["idx_books_author", "books(author_id)", "фильтрация по автору и JOIN authors"],
              ["idx_loans_reader", "loans(reader_id, returned_at)", "GET /readers/{id}/loans, поиск активных выдач"],
              ["idx_loans_book", "loans(book_id, returned_at)", "проверка активных выдач конкретной книги"],
              ["idx_book_tags_tag", "book_tags(tag_id)", "обратная выборка книг по тегу (связь N:M)"],
          ],
          widths=[4.2, 4.6, 8.2])

    body(document, "Назначение индексов проверено планами выполнения (EXPLAIN QUERY PLAN):")
    listing(document, """EXPLAIN QUERY PLAN SELECT * FROM loans WHERE reader_id = 2 AND returned_at IS NULL;
   SEARCH loans USING INDEX idx_loans_reader (reader_id=? AND returned_at=?)

EXPLAIN QUERY PLAN SELECT * FROM books WHERE title = 'Война и мир';
   SEARCH books USING INDEX idx_books_title (title=?)""", "Листинг 11.2 — Использование индексов в плане запроса")

    body(document, "Индекс не создаётся на каждом столбце: каждая дополнительная структура увеличивает размер "
                   "БД и замедляет INSERT, UPDATE и DELETE, так как её нужно поддерживать в актуальном "
                   "состоянии. Кроме того, индекс не используется, если условие начинается с шаблона "
                   "LIKE '%текст%' или если столбец имеет низкую селективность (например, поле с двумя "
                   "значениями). Поэтому индексы созданы только под реальные запросы приложения: поиск по "
                   "названию, фильтрация по автору и выборка выдач по читателю и книге.")

    heading(document, "Транзакции")
    body(document, "Бизнес-операция «выдача книги» изменяет две таблицы: в loans добавляется запись о выдаче, "
                   "а в books уменьшается количество свободных экземпляров. Если выполнить только первую "
                   "операцию, читатель получит книгу, которой нет в наличии, поэтому обе операции объединены "
                   "в транзакцию: при ошибке на любом шаге выполняется ROLLBACK.")
    listing(document, """@contextmanager
def transaction(self):
    connection = self.connect()
    try:
        connection.execute("BEGIN IMMEDIATE")
        yield connection
        connection.commit()      # COMMIT — блок завершился без исключения
    except Exception:
        connection.rollback()    # ROLLBACK — любое исключение внутри блока
        raise
    finally:
        connection.close()""", "Листинг 11.3 — Менеджер транзакций (db.py)")

    listing(document, """with self.database.transaction() as connection:
    row = connection.execute(
        "SELECT available_copies FROM books WHERE id = ?", (book_id,)
    ).fetchone()
    if row is None:
        raise NotFoundError(f"Book with id={book_id} not found")
    if row["available_copies"] <= 0:
        raise ConflictError("Book has no available copies")

    loan_id = self.loans.create(connection, book_id, reader_id, taken_at, due_at)
    cursor = connection.execute(
        "UPDATE books SET available_copies = available_copies - 1"
        " WHERE id = ? AND available_copies > 0", (book_id,)
    )
    if cursor.rowcount != 1:
        raise ConflictError("Book copy was taken by another reader")""",
            "Листинг 11.4 — Транзакция выдачи книги (service.py)")

    body(document, "На рисунке 11.4 показаны три эксперимента: успешная выдача (COMMIT), попытка выдать "
                   "книгу, все экземпляры которой уже на руках (ROLLBACK, состояние БД не изменилось), и "
                   "возврат книги (COMMIT, оба UPDATE применились вместе).")
    figure(document, SCREENS / "11-09-transaction.png",
           "Рисунок 11.4 — Транзакции: COMMIT при выдаче, ROLLBACK при отсутствии экземпляров, COMMIT при возврате")

    heading(document, "Интеграция с Python")
    body(document, "Приложение построено по той же схеме, что и в лабораторной работе №10, но хранилище "
                   "в памяти заменено базой данных:")
    listing(document, """HTTP Client
     │ HTTP
     ▼
Routes (routes.py)      разбор запроса, коды ответов
     ▼
Service (service.py)    бизнес-логика, валидация, кэш
     ▼
Repository (repository.py)  SQL-запросы, параметры, транзакции
     ▼
SQLite (library.db)     таблицы, индексы, ограничения""", "Листинг 11.5 — Архитектура приложения")

    body(document, "Все SQL-запросы выполняются только в слое Repository и только с параметрами. Пример "
                   "параметризованного INSERT:")
    listing(document, """query = '''
    INSERT INTO books (title, author_id, genre, year, price,
                       total_copies, available_copies, cover_path)
    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
'''
params = (data["title"], data["author_id"], data["genre"], data["year"],
          data["price"], data["total_copies"], available, data.get("cover_path"))
with self.database.connection() as connection:
    cursor = connection.execute(query, params)""",
            "Листинг 11.6 — Параметризованный запрос (repository.py)")

    body(document, "Список endpoints, переведённых на работу с БД, приведён в таблице 11.4.")
    table_caption(document, "Таблица 11.4 — Endpoints REST API и выполняемые SQL-операции")
    table(document,
          ["Метод", "Endpoint", "Операция SQL"],
          [
              ["GET", "/books", "SELECT с фильтрами и сортировкой"],
              ["GET", "/books/{id}", "SELECT по первичному ключу (через кэш)"],
              ["GET", "/books/{id}/details", "SELECT с JOIN автора и GROUP_CONCAT тегов"],
              ["GET", "/books/search?q=", "SELECT по названию + INCR счётчика в Redis"],
              ["GET", "/books/search-history", "чтение истории поиска из Redis"],
              ["POST", "/books", "INSERT"],
              ["PUT", "/books/{id}", "UPDATE"],
              ["PATCH", "/books/{id}", "UPDATE + инвалидация кэша"],
              ["DELETE", "/books/{id}", "DELETE"],
              ["PUT", "/books/{id}/tags", "INSERT в book_tags (связь N:M)"],
              ["GET", "/loans", "SELECT с JOIN books и readers"],
              ["POST", "/loans", "транзакция: INSERT + UPDATE"],
              ["POST", "/loans/{id}/return", "транзакция: UPDATE + UPDATE"],
              ["GET", "/stats", "GROUP BY, HAVING, COUNT, AVG"],
              ["GET", "/", "страница каталога (HTML-шаблон)"],
              ["GET", "/catalog", "SELECT книг с ФИО автора и ссылкой на обложку CDN"],
          ],
          widths=[2.2, 5.6, 9.2])

    body(document, "Для демонстрации работы API используется консольный клиент client.py с меню "
                   "(рисунок 11.5).")
    figure(document, SCREENS / "11-13-menu.png", "Рисунок 11.5 — Меню консольного клиента REST API", 14.0)

    body(document, "На рисунке 11.6 приведены список книг (GET /books) и чтение книги по идентификатору "
                   "(GET /books/1, первый запрос — промах кэша, второй — попадание).")
    figure(document, SCREENS / "11-05-client-list.png",
           "Рисунок 11.6 — Получение списка книг и чтение объекта по идентификатору")

    body(document, "На рисунке 11.7 показаны создание (POST), полное изменение (PUT), частичное изменение "
                   "(PATCH) и удаление (DELETE) книги. При изменении и удалении ключ кэша инвалидируется.")
    figure(document, SCREENS / "11-06-client-crud.png",
           "Рисунок 11.7 — Создание, изменение и удаление записи через REST API")

    body(document, "На рисунке 11.8 показаны операции выдачи и возврата книги, выполняемые в транзакции: "
                   "после выдачи количество доступных экземпляров уменьшается, после возврата — "
                   "восстанавливается.")
    figure(document, SCREENS / "11-07-client-transaction.png",
           "Рисунок 11.8 — Выдача книги (INSERT + UPDATE) и возврат книги (UPDATE + UPDATE)")

    heading(document, "NoSQL: Redis")
    body(document, "В качестве NoSQL-хранилища выбран Redis. Выбор обоснован характером данных: история "
                   "поиска и кэш отдельных книг — это небольшие значения по ключу с временем жизни, для "
                   "которых не нужны транзакции и JOIN. Redis хранит такие данные в памяти и выполняет "
                   "GET/SET/INCR за доли миллисекунды, тогда как реляционная БД для этой задачи избыточна.")
    body(document, "На рабочей машине сервер Redis не установлен (официальной сборки для Windows нет), "
                   "поэтому развёрнут локальный Redis-совместимый сервер mini_redis.py, реализующий протокол "
                   "RESP и команды PING, SET (с EX), GET, DEL, EXPIRE, TTL, INCR, KEYS, DBSIZE, FLUSHDB. "
                   "Приложение работает через штатный клиент redis-py, поэтому переход на промышленный "
                   "Redis сводится к изменению host и port в config.py.")

    body(document, "Структура ключей и время жизни приведены в таблице 11.5.")
    table_caption(document, "Таблица 11.5 — Структура ключей Redis и время жизни записей")
    table(document,
          ["Ключ", "Тип данных", "TTL", "Назначение"],
          [
              ["library:book:{id}", "строка (JSON книги)", "60 c", "кэш чтения книги по идентификатору"],
              ["library:search:count:{q}", "счётчик (INCR)", "3600 c", "история и частота поисковых запросов"],
              ["library:demo:*", "строка/счётчик", "2–90 c", "демонстрация команд SET/GET/DEL/EXPIRE/INCR"],
          ],
          widths=[6.0, 4.2, 2.2, 4.6])

    body(document, "Реализованы пять операций Redis: SET (запись значения с TTL), GET (чтение), DEL "
                   "(удаление ключа), EXPIRE (установка времени жизни) и INCR (инкремент счётчика поиска). "
                   "Чтение книги реализовано по паттерну cache-aside:")
    listing(document, """def get_book(self, book_id: int) -> Book:
    key = f"book:{book_id}"
    cached = self.cache.get_product_json(key)      # 1. проверяем Redis
    if cached is not None:
        return Book(**json.loads(cached))          # попадание: ответ из кэша

    book = self.books.get_by_id(book_id)           # 2. промах: SELECT из БД
    if book is None:
        raise NotFoundError(f"Book with id={book_id} not found")

    self.cache.set(key, json.dumps(book.to_dict(), ensure_ascii=False))  # 3. кладём в кэш
    return book""", "Листинг 11.7 — Cache-aside при чтении книги (service.py)")

    body(document, "Инвалидация выполняется при любом изменении данных: после POST, PUT, PATCH, DELETE и "
                   "операций выдачи/возврата ключ book:{id} удаляется, поэтому следующий GET читает свежие "
                   "данные из БД. TTL ограничивает время жизни записи: даже без явной инвалидации "
                   "устаревшее значение исчезнет через 60 секунд, что важно для изменяемых данных (цена, "
                   "количество экземпляров).")
    body(document, "Результаты эксперимента с Redis приведены на рисунке 11.9: показаны базовые команды, "
                   "промах и попадание кэша (2,26 мс против 0,14 мс), инвалидация после PATCH, истечение TTL "
                   "и история поиска.")
    figure(document, SCREENS / "11-10-redis.png",
           "Рисунок 11.9 — Работа Redis: команды, cache-aside (MISS/HIT), инвалидация, TTL, история поиска")

    body(document, "На рисунке 11.10 показано, что REST-запрос поиска увеличивает счётчик в Redis, а "
                   "GET /books/search-history возвращает историю запросов и статистику кэша.")
    figure(document, SCREENS / "11-08-client-redis.png",
           "Рисунок 11.10 — Поиск книги, история поиска и статистика кэша через REST API")

    heading(document, "CDN")
    body(document, "В качестве CDN-контента используются обложки книг и файлы оформления каталога. "
                   "Поскольку внешний CDN в лабораторной работе недоступен, применена разрешённая "
                   "методичкой локальная имитация: origin-сервер (порт 8080) отдаёт файлы из каталога static, "
                   "а edge-сервер CDN (порт 8081) кэширует их и добавляет заголовки X-Cache и Age.")
    body(document, "Размещённые ресурсы и политика кэширования приведены в таблице 11.6.")
    table_caption(document, "Таблица 11.6 — Статические ресурсы, размещённые в CDN")
    table(document,
          ["Ресурс", "URL origin", "URL CDN", "Cache-Control"],
          [
              ["изображения (12 обложек)", "/images/cover-01.png … cover-12.png",
               "http://127.0.0.1:8081/images/cover-01.png", "public, max-age=86400"],
              ["стили", "/css/style.css", "http://127.0.0.1:8081/css/style.css", "public, max-age=3600"],
              ["скрипт каталога", "/js/app.js", "http://127.0.0.1:8081/js/app.js", "public, max-age=3600"],
          ],
          widths=[3.6, 4.6, 5.6, 3.2])

    body(document, "Для каждого ресурса исследованы URL исходного и CDN-ресурса, HTTP-заголовки "
                   "(Content-Type, Content-Length, Cache-Control, Age, X-Cache, X-Served-By), время "
                   "получения и размер ответа, а также разница между первым и последующим запросами "
                   "(рисунок 11.11).")
    figure(document, SCREENS / "11-11-cdn.png",
           "Рисунок 11.11 — Заголовки ответа при прямой загрузке с origin и при загрузке через CDN")

    body(document, "Первый запрос к edge-серверу возвращает X-Cache: MISS и Age: 0 — ресурс берётся с origin и "
                   "сохраняется на edge. Повторный запрос возвращает X-Cache: HIT и Age: 1 — ресурс отдан из "
                   "кэша CDN без обращения к origin. Заголовок Cache-Control: public, max-age=86400 означает, "
                   "что обложка может кэшироваться 24 часа, а стили и скрипты — 1 час (max-age=3600). "
                   "Состояние кэша edge-сервера показано на рисунке 11.12.")
    figure(document, SCREENS / "11-12-cdn-cache.png",
           "Рисунок 11.12 — Состояние кэша edge-сервера: объекты, Age, max-age, количество HIT и MISS")

    body(document, "Принципиальное различие трёх уровней хранения в выполненной работе: цена и количество "
                   "экземпляров книги — изменяемые данные, они хранятся в БД и кэшируются не более 60 секунд; "
                   "история поиска — временные данные, она живёт в Redis; обложка книги — статический файл, "
                   "он отдаётся через CDN и кэшируется на сутки. CDN не заменяет БД: он не умеет "
                   "фильтровать, объединять и изменять данные.")

    heading(document, "Веб-страница каталога")
    body(document, "REST API возвращает данные в формате JSON, поэтому для наглядной демонстрации работы "
                   "приложения добавлена витрина каталога: маршрут GET / отдаёт шаблон templates/index.html, "
                   "а новый endpoint GET /catalog возвращает книги вместе с ФИО автора и ссылкой на обложку. "
                   "Стили и скрипт страницы подключаются с edge-сервера CDN, поэтому открытая в браузере "
                   "страница одновременно показывает работу всех трёх уровней хранения: данные о книгах "
                   "приходят из SQLite, обложки — из CDN, а чтение книги по идентификатору проходит через кэш "
                   "Redis.")
    body(document, "На странице реализованы поиск по названию (параметр q запроса GET /catalog), фильтр "
                   "«только доступные» (available=true) и сортировка по названию, цене, году или автору. "
                   "Карточки книг показывают обложку, название, автора, жанр, год, цену и количество "
                   "доступных экземпляров. При выборе карточки открывается панель с данными запроса "
                   "GET /books/{id}/details: автор, теги из связи N:M и наличие. Кнопка «Оформить выдачу» "
                   "вызывает POST /loans, то есть ту же транзакцию, что и консольный клиент: количество "
                   "экземпляров на карточке сразу уменьшается, потому что после COMMIT кэш книги "
                   "инвалидируется.")
    figure(document, SCREENS / "11-15-catalog.png",
           "Рисунок 11.13 — Страница каталога: данные из SQLite, обложки из CDN, выдача книги через POST /loans")
    body(document, "Страница не заменяет REST API и не содержит собственной логики работы с данными: она "
                   "обращается к тем же endpoints, что и консольный клиент, а вся бизнес-логика остаётся в "
                   "слое Service, поэтому поведение приложения одинаково для браузера и для HTTP-клиента.")

    heading(document, "Сравнение технологий")
    body(document, "Сравнительная характеристика использованных систем хранения приведена в таблице 11.7.")
    table_caption(document, "Таблица 11.7 — Сравнительная характеристика систем хранения")
    table(document,
          ["Характеристика", "PostgreSQL / SQLite", "MongoDB", "Redis", "CDN"],
          [
              ["Основное назначение", "основное хранилище данных приложения", "документное хранилище",
               "кэш, счётчики, сессии", "доставка статического контента"],
              ["Структура", "таблицы, строки, столбцы", "коллекции документов (BSON)", "ключ → значение",
               "файлы по URL"],
              ["Хранение", "на диске, с журналом транзакций", "на диске", "в оперативной памяти",
               "на edge-серверах"],
              ["Запросы", "SQL: JOIN, GROUP BY, подзапросы", "find/aggregate", "GET/SET/INCR по ключу",
               "HTTP GET файла"],
              ["Транзакции", "полноценные, ACID", "много­документные транзакции", "атомарные команды, без "
               "ACID-набора", "нет"],
              ["Кэширование", "нет встроенного кэша чтения", "кэш в памяти (WiredTiger)",
               "сам является кэшем, TTL", "кэш с Cache-Control и Age"],
              ["Масштабирование", "репликация, шардирование", "шардирование, репликация", "кластер, реплики",
               "сеть edge-серверов"],
              ["Типичные данные", "заказы, книги, читатели, платежи", "профили, каталоги с гибкой схемой",
               "сессии, счётчики, история поиска", "изображения, CSS, JS, видео"],
              ["Когда использовать", "целостность и связи важнее скорости записи",
               "структура документов меняется", "нужны очень быстрые чтения по ключу",
               "много статики и географически распределённые пользователи"],
          ],
          widths=[3.4, 3.6, 3.2, 3.4, 3.4], font_size=10)

    heading(document, "Тестирование")
    body(document, "Написано 51 автоматический тест на pytest. Тесты не требуют запущенных серверов: "
                   "используется временная БД, создаваемая из schema.sql и seed.sql, и тестовый кэш с тем же "
                   "интерфейсом, что у RedisCache. Состав тестов приведён в таблице 11.8.")
    table_caption(document, "Таблица 11.8 — Состав автоматических тестов")
    table(document,
          ["Файл", "Тестов", "Что проверяется"],
          [
              ["tests/test_repository.py", "15", "SQL-запросы, фильтры, сортировка, JOIN, связь N:M, "
                                                 "GROUP BY/HAVING, ограничения внешних ключей, "
                                                 "попытка SQL-инъекции"],
              ["tests/test_service.py", "13", "бизнес-логика, cache-aside, инвалидация, валидация полей, "
                                              "поиск и история поиска, статистика"],
              ["tests/test_transaction.py", "9", "выдача и возврат книги, откат при отсутствии экземпляров, "
                                                 "повторный возврат, инвалидация кэша"],
              ["tests/test_api.py", "14", "все endpoints REST API, коды 201/204/400/404/409, "
                                          "чтение и запись через HTTP"],
          ],
          widths=[5.2, 1.8, 9.5])

    body(document, "Результат прогона тестов приведён на рисунке 11.14: все 51 тест пройдены.")
    figure(document, SCREENS / "11-04-pytest.png", "Рисунок 11.14 — Результат выполнения тестов pytest")

    heading(document, "Вывод")
    body(document, "В ходе лабораторной работы разработана система хранения данных для предметной области "
                   "«Библиотека» (вариант 2). Спроектирована реляционная схема из шести таблиц с первичными "
                   "и внешними ключами, связями 1:N и N:M через промежуточную таблицу book_tags; схема "
                   "приведена к 3НФ. Подготовлены schema.sql, seed.sql и 15 SQL-запросов, охватывающих "
                   "SELECT, WHERE, ORDER BY, JOIN, LEFT JOIN, GROUP BY, HAVING, агрегатные функции, "
                   "подзапросы, INSERT, UPDATE и DELETE. Создано пять индексов, назначение которых "
                   "подтверждено планами выполнения запросов.")
    body(document, "Приложение из лабораторной работы №10 переведено на работу с БД: все endpoints обращаются "
                   "к данным через слой Repository, SQL-запросы параметризованы, что исключает SQL-инъекции. "
                   "Реализована транзакционная бизнес-операция выдачи книги: добавление записи о выдаче и "
                   "уменьшение количества свободных экземпляров выполняются как единое целое, а при ошибке "
                   "выполняется ROLLBACK, и состояние БД не меняется.")
    body(document, "Дополнительно реализованы Redis-кэш по паттерну cache-aside с TTL 60 секунд, "
                   "автоматической инвалидацией при изменении данных и историей поиска на счётчиках INCR, а "
                   "также раздача статических ресурсов через CDN с заголовками Cache-Control, Age и X-Cache. "
                   "Проведённые эксперименты показали разницу между уровнями хранения: изменяемые данные "
                   "(цена, наличие) должны поступать из БД, временные (история поиска) — из Redis, "
                   "статические (обложки) — из CDN. Написано 51 автоматический тест, все тесты проходят.")
    body(document, "Главный результат работы — понимание того, что современное приложение использует "
                   "несколько систем хранения одновременно, и выбор конкретной системы определяется "
                   "характером данных, требованиями к целостности и скорости доступа.")

    return document


if __name__ == "__main__":
    OUT.parent.mkdir(parents=True, exist_ok=True)
    document = build()
    document.save(OUT)
    print(f"Отчёт сохранён: {OUT}")
