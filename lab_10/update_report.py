"""Добавить в отчёт ЛР №10 раздел «Веб-страница каталога» со скриншотом.

Используется копирование XML готовых параграфов, чтобы сохранить оформление.
"""

import copy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm

REPORT = Path(r"C:\elkin\Отчёты\Практика 10 Василенко И.А.docx")
SCREENSHOT = Path(r"C:\elkin\lab_10\screens\10-09-catalog.png")

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def find_paragraph(document: Document, text: str):
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    return None


def replace_text_in_element(paragraph_xml, new_text: str) -> None:
    """Оставить только первый run, заменить в нём текст."""
    runs = paragraph_xml.findall(f"{{{W}}}r")
    for run in runs[1:]:
        paragraph_xml.remove(run)
    if runs:
        first = runs[0]
        texts = first.findall(f"{{{W}}}t")
        for text_node in texts[1:]:
            first.remove(text_node)
        if texts:
            texts[0].set(f"{{{W}}}space", "preserve")
            texts[0].text = new_text
        else:
            node = first.makeelement(f"{{{W}}}t", {f"{{{W}}}space": "preserve"})
            node.text = new_text
            first.append(node)
    else:
        run = paragraph_xml.makeelement(f"{{{W}}}r", {})
        node = paragraph_xml.makeelement(f"{{{W}}}t", {f"{{{W}}}space": "preserve"})
        node.text = new_text
        run.append(node)
        paragraph_xml.append(run)


def main() -> None:
    document = Document(REPORT)

    heading_anchor = find_paragraph(document, "Вывод")
    if heading_anchor is None:
        raise RuntimeError("Раздел «Вывод» не найден")

    caption_sample = find_paragraph(document, "Рисунок 10.1 — Результат выполнения автоматических тестов pytest")
    if caption_sample is None:
        caption_sample = heading_anchor
        for paragraph in document.paragraphs:
            if paragraph.text.strip().startswith("Рисунок 10.1"):
                caption_sample = paragraph
                break

    body_sample = None
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        style_name = paragraph.style.name if paragraph.style else ""
        if len(text) > 200 and not style_name.startswith("Heading"):
            body_sample = paragraph
            break
    if body_sample is None:
        body_sample = heading_anchor

    heading_template = heading_anchor._p

    elements = []

    heading = copy.deepcopy(heading_template)
    replace_text_in_element(heading, "Веб-страница каталога")
    elements.append(heading)

    body_text_1 = (
        "REST API возвращает данные в формате JSON, и HTTP-клиентом может быть не только "
        "программа на Python, но и браузер. Чтобы продемонстрировать это, поверх разработанного "
        "API добавлена витрина каталога: маршрут GET / отдаёт шаблон templates/index.html, который "
        "загружает данные через те же endpoints GET /books и GET /books/{id}, что и консольная "
        "программа client.py. Таким образом, страница не содержит собственной логики работы с "
        "данными — вся бизнес-логика по-прежнему находится в слое Service, а хранилище — в "
        "Repository."
    )
    body_1 = copy.deepcopy(body_sample._p)
    replace_text_in_element(body_1, body_text_1)
    elements.append(body_1)

    body_text_2 = (
        "На странице реализованы поиск по названию или автору, фильтр «только доступные» "
        "(query-параметр available=true) и сортировка по названию, автору, цене или году "
        "(sort и order). При выборе карточки браузер выполняет запрос GET /books/{id} и "
        "показывает детали книги. Это подтверждает важный вывод лабораторной работы: любой "
        "HTTP-клиент — консольная программа, тест или браузер — работает с разработанным API "
        "одинаково, с одним набором endpoints, методов и кодов состояния."
    )
    body_2 = copy.deepcopy(body_sample._p)
    replace_text_in_element(body_2, body_text_2)
    elements.append(body_2)

    picture_paragraph = document.add_paragraph()
    picture_run = picture_paragraph.add_run()
    picture_run.add_picture(str(SCREENSHOT), width=Cm(15.5))
    picture_paragraph.alignment = 1  # CENTER
    elements.append(picture_paragraph._p)

    caption = copy.deepcopy(caption_sample._p)
    replace_text_in_element(caption, "Рисунок 10.9 — Страница каталога, работающая поверх REST API (GET /)")
    elements.append(caption)

    anchor = heading_anchor._p
    for element in elements:
        anchor.addprevious(element)

    document.save(REPORT)
    print(f"Готово: {REPORT}")


if __name__ == "__main__":
    main()